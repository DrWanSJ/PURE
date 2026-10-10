"""Read-only S12-S16/S27 and complete relevant original article pages."""
import sys,json,csv,warnings,argparse,hashlib
from pathlib import Path
from collections import defaultdict
from fractions import Fraction as F
sys.dont_write_bytecode=True
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'docs/reduction/pathways'
def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def main():
    from openpyxl import load_workbook
    from pypdf import PdfReader
    p=argparse.ArgumentParser();p.add_argument('--output-dir',type=Path,default=OUT);args=p.parse_args()
    inv=json.loads((args.output_dir/'phase_b1_3_source_scope.json').read_text(encoding='utf-8'));data=json.loads((args.output_dir/'phase_b1_3_witnesses.json').read_text(encoding='utf-8'))
    folder=ROOT/'references/PNAS2017_Matsuura/raw';paper=next(folder.glob('*.pdf'));files=[folder/f'pnas.1615351114.sd{i:02d}.xlsx' for i in [12,13,14,15,16,27]]+[paper]
    before={f.relative_to(ROOT).as_posix():digest(f) for f in files};mapping=defaultdict(list);rows=[];notices=[]
    for r,q in inv['reactions'].items():
        for m in q['original_subsystem_memberships']:mapping[Path(m['source_file']).stem,m['local_reaction_id']].append(r)
    for i in [12,13,14,15,16]:
        path=folder/f'pnas.1615351114.sd{i:02d}.xlsx'
        with warnings.catch_warnings(record=True) as ws:
            book=load_workbook(path,read_only=True,data_only=False);model=book['Model']['A1'].value
            for n,row in enumerate(book['Reactions'].values,1):
                if n==1 or not row[1]:continue
                ids=mapping[model,row[1]];assert len(ids)==1,(model,row[1],ids);r=ids[0];q=inv['reactions'][r]
                def terms(x):return {s:str(n) for s,n in sorted(__import__('collections').Counter(t.strip() for t in str(x or '').split(',') if t.strip()).items())}
                agree=terms(row[5])==q['reactants'] and terms(row[6])==q['products']
                rows.append({'file':path.relative_to(ROOT).as_posix(),'sheet':'Reactions','row':n,'local_id':row[1],'canonical_id':r,'equation':q['equation'],'canonical_sides_agree':agree,'authority':'AUTHOR_SUBSYSTEM_WORKBOOK_CONTEXT'})
            book.close();notices.extend({'file':path.name,'warning':str(x.message)} for x in ws)
    ids={e['reaction_id'] for w in data['witnesses'] for e in w['reaction_occurrences'] if e['source_segment']=='B1_3'}
    ids|={r for a in list(ids) for r in inv['reactions'][a]['reverse_reaction_ids']}
    species={s for r in ids for side in ['reactants','products'] for s in inv['reactions'][r][side]}
    legends=[];parameters=[];refs=[];initial={}
    path=folder/'pnas.1615351114.sd27.xlsx'
    with warnings.catch_warnings(record=True) as ws:
        book=load_workbook(path,read_only=True,data_only=False)
        for n,row in enumerate(book['initial concentrations'].values,1):
            if row[2] in species:legends.append({'species_id':row[2],'row':n,'sheet':'initial concentrations','definition':row[4],'initial_uM':row[3],'authority':'AUTHOR_S27_DEFINITION','evidence_status':'EXTRACTED'})
        for n,row in enumerate(book['reactions and parameters'].values,1):
            if row[0] in ids:parameters.append({'reaction_id':row[0],'row':n,'sheet':'reactions and parameters','reactants_cell':row[1],'products_cell':row[2],'value':row[3],'unit':row[4],'reverse_cell':row[5],'note':row[6],'reference_number':row[7],'parameter_tag':row[8],'canonical_parameter_agrees':F(str(row[3]))==F(inv['reactions'][row[0]]['reference_parameter'])})
        refnumbers={x['reference_number'] for x in parameters}
        for n,row in enumerate(book['references for parameters'].values,1):
            if row[0] in refnumbers:refs.append({'row':n,'number':row[0],'reference':row[1],'PMID':row[2]})
        book.close();notices.extend({'file':path.name,'warning':str(x.message)} for x in ws)
    csvpath=ROOT/'models/pnas2017_full_reference/original/simulate/Simulate_fMGG_synthesis/dat/fMGG_synthesis_initial_values.csv'
    with csvpath.open(encoding='utf-8-sig',newline='') as f:
        for row in csv.DictReader(f):
            if row['Name'] in species:initial[row['Name']]=row['Value']
    pdf=PdfReader(paper);pages={n:pdf.pages[n-1].extract_text() for n in [2,7,8]}
    assert '241' in pages[2] and '968' in pages[2]
    locators=[{'page':n,'keyword':k,'present':k.lower() in text.lower()} for n,text in pages.items() for k in ['termination','recycling','RF3','EF-G']]
    changed=[s for s,h in before.items() if digest(ROOT/s)!=h];assert not changed
    assert legends and parameters,'EMPTY_S27_EVIDENCE'
    initial_comparisons=[{'species_id':x['species_id'],'S27_uM':str(x['initial_uM']),'operational_CSV_value':initial.get(x['species_id']),'agrees':F(str(x['initial_uM']))==F(initial[x['species_id']]) if x['species_id'] in initial else None,'authority':'OPERATIONAL_CSV_FOR_OPERATION; S27_RETAINED_AS_ORIGINAL_DEFINITION'} for x in legends]
    report={'status':'PASS' if all(r['canonical_sides_agree'] for r in rows) and all(r['canonical_parameter_agrees'] for r in parameters) else 'FAIL','source_hashes':before,'unchanged_original_files':len(before),'subsystem_reaction_rows':rows,'S27_definitions':legends,'S27_parameters':parameters,'parameter_references':refs,'author_operational_initial_values':initial,'initial_concentration_comparison':initial_comparisons,'paper':{'path':paper.relative_to(ROOT).as_posix(),'doi':'10.1073/pnas.1615351114','full_pages_read':[2,7,8],'keyword_locators':locators,'evidence_status':'EXTRACTED_LIMITED_CONTEXT','claims':'Article establishes combined model architecture; exact B1-3 steps are established by canonical SBML/S12-S16/S27, not by a physiological inference from paper figures.'},'reader_warnings':notices,'complete_standalone_SI_text_pdf':'NOT_ESTABLISHED_LOCALLY','composition_certificate':'NOT_ESTABLISHED','source_state_contract':'Opaque finite carrier lineage projected from original source IDs; molecular interpretation remains qualified.'}
    (args.output_dir/'phase_b1_3_original_source_evidence.json').write_bytes((json.dumps(report,ensure_ascii=False,indent=2,sort_keys=True)+'\n').encode())
    lines=['# B1-3 原始来源与解释边界','', '规范SBML和作者CSV最高优先；本轮只读原始S12-S16和S27，所有来源原始字节保持。源反应结构EXTRACTED；完整分子组成及真实生理顺序未认证。','', '## 原始文件定位','']
    for path,h in before.items():lines.append('- `'+path+'` SHA-256 `'+h+'`。')
    lines+=['','原论文完整读取页2/7/8（doi 10.1073/pnas.1615351114），支持模型架构背景。具体终止/回收以原始子系统和规范SBML逐方向核对，不能从文章图推断遗漏步骤。独立完整SI文本/PDF本地仍未建立。','', '## S27作者源物种定义','', '| 源物种 | S27行 | 作者定义（摘录原始数据） | 初始uM |','|---|---:|---|---:|']
    for r in legends:lines.append('| `'+r['species_id']+'` | '+str(r['row'])+' | '+str(r['definition']).replace('|','/')+' | '+str(r['initial_uM'])+' |')
    lines+=['','T_pre是项目缩写；作者把RF-free入口称为elongation complex，RF结合态称为pre-termination complex。RF3/EF-G的GDP/GTP/PO4状态保留为原始独立物种。S27和CSV的初始条件不同不能自动用有限witness lot替代作者浓度；操作CSV原值见机器附录。','', '初始浓度原始差异：`'+json.dumps([x for x in initial_comparisons if x['agrees'] is False],ensure_ascii=False)+'`。不调和或改写来源。','', '## S27参数、假设与原始逆向列','', '| 源ID | S27行 | 参数 | 说明 | 作者逆向列 |','|---|---:|---:|---|---|']
    for r in parameters:lines.append('| '+r['reaction_id']+' | '+str(r['row'])+' | '+str(r['value'])+' | '+str(r['note'] or '')+' | '+str(r['reverse_cell'] or '')+' |')
    lines+=['','反向关系独立由规范方程双边精确交换重算；作者逆向列仅保留原始语义。参数正值仅代表方向可用，不证明主导通量。完整S12-S16行级映射及零参数/降解记录见JSON和source_scope。','', f'本轮核对{len(rows)}条五子系统工作簿反应行、{len(parameters)}条S27参数、{len(legends)}条S27定义；不保存来源工作簿。完整元素、核苷酸基团、电荷/渗透守恒、H2O/H+显式化均未证明。','']
    (args.output_dir/'phase_b1_3_original_source_evidence.md').write_bytes(('\n'.join(lines)+'\n').encode())
    print(json.dumps({'status':report['status'],'subsystem_rows':len(rows),'S27_parameters':len(parameters),'S27_definitions':len(legends),'unchanged_files':len(before)}));return 0 if report['status']=='PASS' else 1
if __name__=='__main__':raise SystemExit(main())
