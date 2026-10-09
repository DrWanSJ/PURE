#!/usr/bin/env python3
"""Read-only local paper/S05-S11/S27 interpretation evidence; never save sources."""
from collections import defaultdict
from fractions import Fraction
from pathlib import Path
import argparse
import hashlib
import json
import sys
import warnings

sys.dont_write_bytecode=True
sys.stdout.reconfigure(encoding='utf-8')
ROOT=Path(__file__).resolve().parents[3];OUT=ROOT/'docs/reduction/pathways'

def digest(p):return hashlib.sha256(p.read_bytes()).hexdigest()
def save(p,d):p.write_bytes((json.dumps(d,sort_keys=True,indent=2,ensure_ascii=False)+'\n').encode())

def main():
    from openpyxl import load_workbook
    from pypdf import PdfReader
    p=argparse.ArgumentParser(description=__doc__);p.add_argument('--output-dir',type=Path,default=OUT);args=p.parse_args()
    scope=json.loads((args.output_dir/'phase_b1_2_source_scope.json').read_text(encoding='utf-8'))
    data=json.loads((args.output_dir/'phase_b1_2_witnesses.json').read_text(encoding='utf-8'))
    w=next(w for w in data['witnesses'] if w['witness_id']=='W3')
    ids={e['reaction_id'] for e in w['reaction_occurrences']}|{'re0000000001','re0000000002'}
    ids|={r for t in list(ids) for r in scope['reactions'][t]['reverse_reaction_ids']}
    ids|=set(scope['boundary_incidence'][w['target_source_state']]['consumers'])
    species={s for e in w['reaction_occurrences'] for side in ('inputs','outputs') for s in e[side]}|{'elRS70SAGGU0002_fMettRNAfMetCAU'}
    folder=ROOT/'references/PNAS2017_Matsuura/raw';source_rows=[];legends=[];mapped=[];notices=[]
    before={p.relative_to(ROOT).as_posix():digest(p) for p in [*folder.glob('pnas.1615351114.sd0[5-9].xlsx'),*folder.glob('pnas.1615351114.sd1[01].xlsx'),folder/'pnas.1615351114.sd27.xlsx',*folder.glob('*.pdf')]}
    reverse_lookup=defaultdict(list)
    for r,q in scope['reactions'].items():
        for m in q['original_subsystem_memberships']:reverse_lookup[Path(m['source_file']).stem,m['local_reaction_id']].append(r)
    for i in range(5,12):
        file=folder/f'pnas.1615351114.sd{i:02}.xlsx'
        with warnings.catch_warnings(record=True) as seen:
            book=load_workbook(file,read_only=True,data_only=False)
            name=book['Model']['A1'].value
            for rownum,row in enumerate(book['Reactions'].values,1):
                if rownum==1:continue
                local=row[1];combined=reverse_lookup[name,local]
                if not local:continue
                if len(combined)!=1:raise ValueError('SUBSYSTEM_XLSX_MAPPING_AMBIGUOUS:'+str((name,local,combined)))
                r=combined[0];q=scope['reactions'][r]
                def terms(raw):
                    acc=defaultdict(Fraction)
                    for s in str(raw or '').split(','):
                        if s:acc[s.strip()]+=1
                    return {s:str(v) for s,v in sorted(acc.items())}
                agree=terms(row[5])==q['reactants'] and terms(row[6])==q['products']
                mapped.append({'file':file.relative_to(ROOT).as_posix(),'sheet':'Reactions','row':rownum,
                    'local_model_id':name,'local_reaction_id':local,'canonical_reaction_id':r,
                    'source_reaction_type':row[0],'source_reactants':row[5],'source_products':row[6],
                    'agrees_with_canonical_equation':agree,'authority':'SUPPLEMENT_CONTEXT_CANONICAL_OVERRIDES',
                    'evidence_status':'EXTRACTED'})
            book.close();notices += [{'file':file.name,'warning':str(x.message),'sources_saved':False} for x in seen]
    s27=folder/'pnas.1615351114.sd27.xlsx'
    with warnings.catch_warnings(record=True) as seen:
        book=load_workbook(s27,read_only=True,data_only=False)
        for n,row in enumerate(book['initial concentrations'].values,1):
            if row[2] in species:
                legends.append({'file':s27.relative_to(ROOT).as_posix(),'sheet':'initial concentrations','row':n,
                    'species_id':row[2],'author_initial_concentration_uM':row[3],'author_legend':row[4],
                    'evidence_status':'EXTRACTED_AUTHOR_DEFINITION','physiological_identity':'UNRESOLVED',
                    'full_molecular_composition_certificate':False})
        for n,row in enumerate(book['reactions and parameters'].values,1):
            if row[0] in ids:
                q=scope['reactions'][row[0]]
                source_rows.append({'file':s27.relative_to(ROOT).as_posix(),'sheet':'reactions and parameters','row':n,
                    'reaction_id':row[0],'reactants_cell':row[1],'products_cell':row[2],'parameter_cell':row[3],
                    'unit_cell':row[4],'reverse_cell':row[5],'note_cell':row[6],'reference_number_cell':row[7],
                    'parameter_tag_cell':row[8],'author_csv_agrees':Fraction(str(row[3]))==Fraction(q['reference_parameter']),
                    'canonical_equation':q['equation'],'authority':'SUPPLEMENT_CONTEXT_CANONICAL_AND_AUTHOR_CSV_OVERRIDE'})
        references=[{'row':n,'reference_number':row[0],'reference':row[1],'PMID':row[2]} for n,row in enumerate(book['references for parameters'].values,1) if row[0] in {1,2}]
        book.close();notices += [{'file':s27.name,'warning':str(x.message),'sources_saved':False} for x in seen]
    article=next(folder.glob('*.pdf'));pdf=PdfReader(article)
    # Inspect full relevant pages, retain only short locators and paraphrases.
    pages={i+1:pdf.pages[i].extract_text() for i in (1,6,7)}
    assert all(k in pages[2] for k in ('241','968','elongation'))
    assert 'reaction 22' in pages[7] and 'EF-G' in pages[7]
    paper={'file':article.relative_to(ROOT).as_posix(),'sha256':digest(article),'pages_read':[2,7,8],
        'doi':'10.1073/pnas.1615351114','evidence_status':'EXTRACTED_LIMITED_ORIGINAL_ARTICLE_CONTEXT',
        'supported_paraphrases':[{'page':2,'claim':'The fMGG model comprises functional modules, 241 components and 968 reactions.'},
            {'page':7,'claim':'The authors describe deactivating translating-ribosome EF-G GTP hydrolysis by setting reaction 22 to zero.'},
            {'page':8,'claim':'Authors built and combined subsystem models into the full SBML model.'}],
        'limits':'No experimental confirmation of every virtual source-state step or complete composite molecular inventory is established here.'}
    changed=[p for p,h in before.items() if digest(ROOT/p)!=h]
    if changed:raise ValueError('SOURCE_FILE_CHANGED:'+str(changed))
    report={'status':'PASS','source_hashes':before,'source_files_unchanged':len(before),'paper':paper,
        'seven_elongation_xlsx_reaction_rows':mapped,'local_xlsx_rows_checked':len(mapped),
        'local_xlsx_equation_discrepancies':[x for x in mapped if not x['agrees_with_canonical_equation']],
        's27_legends':legends,'s27_reaction_rows':source_rows,'s27_reference_numbers_context_only':references,
        's27_parameter_discrepancies':[x for x in source_rows if not x['author_csv_agrees']],
        'reader_warnings':notices,'complete_standalone_SI_text_pdf':'NOT_ESTABLISHED_LOCALLY',
        'boundary_naming_issue':'S27 calls factor-free T_pre an elongation complex at stop codon; RF1/RF2-bound states are named pre-termination complexes. T_pre is a proposed structural interface label, pending H6.',
        'virtual_state_issue':'S27 explicitly calls E2 and the second-round entry virtual elongation complexes. Exact equations are author-model structure, not complete physiological steps.'}
    save(args.output_dir/'phase_b1_2_original_source_evidence.json',report)
    text=['# B1-2 original-source interpretation appendix','',
          'Read-only inspection of local article pages 2, 7, 8 and original S05-S11/S27 workbooks. Canonical SBML/author CSV retains priority. Sources were never saved or altered.','',
          '## Author definitions and scientific boundaries','',report['virtual_state_issue'],'',report['boundary_naming_issue'],'',
          'The author-defined phrase and molecular interpretation are distinct: author definitions are EXTRACTED, full composition and physiological identity remain INFERRED/UNRESOLVED. The proposed T_pre label is not substituted by a release-factor-bound endpoint.','',
          '## Relevant S27 source definitions','', '| Source species | S27 row | Author model definition | Initial uM |','|---|---:|---|---:|']
    for row in legends:text.append(f"| `{row['species_id']}` | {row['row']} | {row['author_legend']} | {row['author_initial_concentration_uM']} |")
    text+=['','## Reaction notes and assumptions','', '| Original ID | S27 row | Author note | Reference no. | Parameter agreement |','|---|---:|---|---|---|']
    for row in source_rows:text.append(f"| {row['reaction_id']} | {row['row']} | {row['note_cell'] or ''} | {row['reference_number_cell']} | {row['author_csv_agrees']} |")
    text+=['','Several source events are explicitly assumed fast (including 0001, 0016, 0018, 0025, 0068 and their round-two counterparts). This is author model provenance, not measured path dominance or elapsed-time evidence. The physiological ordering of early initiator/inter-round tRNA release and peptidyl transfer remains a Human Review issue.','',
           '## Original article and supplement coverage','',
           'Article p.7 provides limited direct context for reaction 22 as EF-G GTP hydrolysis on translating ribosomes. S27 defines bound GDP/GTP complexes and translocation states. This supports interpretation as author-model steps; no detailed SI text/PDF or experimental certification is invented.','',
           f"Seven original elongation workbook models: {len(mapped)} reaction rows mapped by local model/ID to canonical exact equations; {len(report['local_xlsx_equation_discrepancies'])} equation discrepancies. S27 checked parameter rows: {len(source_rows)}; discrepancies: {len(report['s27_parameter_discrepancies'])}.",'',
           'The workbook reader warns about unsupported extensions. It is used read-only and every inspected source raw hash remains unchanged. The machine appendix records every warning, row locator, mapping and any discrepancy.','']
    (args.output_dir/'phase_b1_2_original_source_evidence.md').write_bytes(('\n'.join(text)+'\n').encode())
    print(json.dumps({'status':'PASS','local_xlsx_rows':len(mapped),'s27_legends':len(legends),'s27_parameter_rows':len(source_rows),'source_discrepancies':len(report['local_xlsx_equation_discrepancies'])+len(report['s27_parameter_discrepancies'])}));return 0

if __name__=='__main__':raise SystemExit(main())
